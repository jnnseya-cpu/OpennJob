"""Discovery, documents, the report and its scheduler, the loopback server and the review commands."""
import json, os, smtplib, threading, unittest, urllib.request, urllib.error
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo
from agent import paths, discovery, report as report_mod
from agent.core import Store, now
from agent.documents import build, PackError
from tests.helpers import DataDir, profile, fixture_job, fixture_adapter


class Discovery(unittest.TestCase):
    def feed(self, description):
        return [{'id': 'feed-1', 'posting_id': '1', 'company': 'Example', 'title': 'Senior Project Manager', 'url': 'https://example.org/1', 'description': description, 'requirements': [], 'country': 'United Kingdom'}]

    def test_changed_description_resets_review_and_failures_are_logged(self):
        with DataDir():
            (paths.data_dir() / 'boards.json').write_text(json.dumps([{'company': 'Example', 'type': 'greenhouse', 'board': 'x'}, {'company': 'Broken', 'type': 'greenhouse', 'board': 'y'}]))
            s = Store(paths.db_path())
            def fetch(board, desc=['A']):
                if board['company'] == 'Broken':
                    raise TimeoutError()
                return self.feed(desc[0])
            with patch.object(discovery, 'fetch', fetch), patch.object(discovery.llm, 'match', side_effect=RuntimeError('no key')):
                discovery.once(s, max_matches=1)
                job = s.jobs()[0]; job.update(requirements=[{'text': 't', 'weight': 100, 'state': 'met', 'evidence_ids': ['E01'], 'hard': False}], matching_reviewed=True, preferences_confirmed=True); s.put_job(job)
                discovery.once(s, max_matches=1)
                self.assertTrue(s.jobs()[0]['matching_reviewed'])  # unchanged description keeps the review
                with patch.object(discovery, 'fetch', lambda b: fetch(b, ['B'])):
                    discovery.once(s, max_matches=1)
            job = s.jobs()[0]
            self.assertEqual(len(s.jobs()), 1)
            self.assertFalse(job['matching_reviewed']); self.assertEqual(job['requirements'], []); self.assertFalse(job['preferences_confirmed'])
            kinds = [r['kind'] for r in s.db.execute('SELECT kind FROM events')]
            self.assertIn('source_failed', kinds); self.assertIn('description_changed', kinds); self.assertIn('matching_failed', kinds)
            s.db.close()

    def test_matching_budget_is_respected(self):
        with DataDir():
            (paths.data_dir() / 'boards.json').write_text(json.dumps([{'company': 'Example', 'type': 'greenhouse', 'board': 'x'}]))
            s = Store(paths.db_path())
            many = [dict(self.feed('d')[0], id=f'j{i}', posting_id=str(i), url=f'https://example.org/{i}') for i in range(5)]
            calls = []
            with patch.object(discovery, 'fetch', lambda b: many), patch.object(discovery.llm, 'match', lambda j, p: calls.append(j['id']) or j):
                discovery.once(s, max_matches=2)
            self.assertEqual(len(calls), 2); s.db.close()


class Documents(unittest.TestCase):
    def test_deterministic_extractive_and_sealed(self):
        import tempfile
        with tempfile.TemporaryDirectory() as a, tempfile.TemporaryDirectory() as b:
            p1, p2 = build(fixture_job(), profile(), a), build(fixture_job(), profile(), b)
            self.assertEqual(p1['document_hashes'], p2['document_hashes'])  # byte-identical files
            self.assertEqual(build(fixture_job(), profile(), a)['integrity'], p1['integrity'])  # same seal when rebuilt in place
            self.assertIn('Led contractor governance on a £40M rail depot upgrade. (fictional)', p1['cv_text'])
            self.assertIn('Example Rail Ltd (fictional) — Construction Manager | 2019-01–2025-12', p1['cv_text'])
            for f in ('cv.pdf', 'cover.pdf', 'cv.docx', 'cover.docx', 'jd_snapshot.txt', 'scorecard.json'):
                self.assertTrue((Path(a) / f).is_file(), f)
            self.assertTrue((Path(a) / 'cv.pdf').read_bytes().startswith(b'%PDF'))
            from docx import Document
            self.assertIn('Alex Example', [x.text for x in Document(str(Path(a) / 'cv.docx')).paragraphs][0])

    def test_refuses_unconfirmed_or_unknown_evidence_and_placeholders(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            j = fixture_job(); j['requirements'][0]['evidence_ids'] = ['E04']  # not confirmed by the applicant
            with self.assertRaises(PackError):
                build(j, profile(), d)
            j['requirements'][0]['evidence_ids'] = ['E99']
            with self.assertRaises(PackError):
                build(j, profile(), d)
            p = profile(); p['employment'][0]['end'] = '[End date required]'
            with self.assertRaises(PackError):
                build(fixture_job(), p, d)
            p = profile(); del p['phone']
            with self.assertRaises(PackError):
                build(fixture_job(), p, d)


class Report(unittest.TestCase):
    def store(self):
        s = Store(':memory:'); j = fixture_job(); s.put_job(j); s.draft(j, {}); s.transition('fixture', 'ready'); s.transition('fixture', 'submitting'); s.transition('fixture', 'uncertain')
        return s

    def test_report_sections(self):
        text = report_mod.report(self.store())
        for part in ('LAST 24 HOURS', 'UNCERTAIN: 1', 'READY means waiting for you to submit', 'RECENT SOURCE FAILURES', 'Local fixture — Construction Manager'):
            self.assertIn(part, text)

    def test_09_00_london_once_per_day_with_dst_and_no_blind_retry(self):
        s = self.store(); sent = []
        london = ZoneInfo('Europe/London')
        # 07:59 UTC on 6 October is 08:59 BST: too early.
        self.assertIsNone(report_mod.tick(s, datetime(2026, 10, 6, 7, 59, tzinfo=ZoneInfo('UTC')), sent.append))
        self.assertEqual(report_mod.tick(s, datetime(2026, 10, 6, 8, 0, tzinfo=ZoneInfo('UTC')), sent.append), 'sent')
        self.assertIsNone(report_mod.tick(s, datetime(2026, 10, 6, 15, 0, tzinfo=london), sent.append))
        # Winter time: 09:00 GMT is 09:00 UTC. A start after 09:00 catches up once.
        self.assertIsNone(report_mod.tick(s, datetime(2026, 12, 1, 8, 59, tzinfo=ZoneInfo('UTC')), sent.append))
        self.assertEqual(report_mod.tick(s, datetime(2026, 12, 1, 13, 0, tzinfo=ZoneInfo('UTC')), sent.append), 'sent')
        self.assertEqual(len(sent), 2)
        def fail(store):
            raise smtplib.SMTPServerDisconnected('dropped')
        self.assertEqual(report_mod.tick(s, datetime(2026, 12, 2, 9, 30, tzinfo=london), fail), 'failed_or_uncertain')
        self.assertIsNone(report_mod.tick(s, datetime(2026, 12, 2, 10, 30, tzinfo=london), fail))  # not retried

    def test_send_needs_smtp_settings(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaises(RuntimeError):
                report_mod.send(self.store())


class Server(unittest.TestCase):
    TOKEN = 'test-token-0123456789-abcdefghij'

    def call(self, method, path, body=None, token=TOKEN):
        req = urllib.request.Request(self.base + path, method=method, data=json.dumps(body).encode() if body is not None else None,
                                     headers={'Authorization': 'Bearer ' + token, 'Content-Type': 'application/json'})
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.load(r)
        except urllib.error.HTTPError as e:
            return e.code, json.load(e)

    def test_pack_attempt_result_flow(self):
        from agent.server import make_server
        with DataDir(), patch.dict(os.environ, {'CAREER_AGENT_TOKEN': self.TOKEN}):
            s = Store(paths.db_path()); j = fixture_job(); s.put_job(j)
            server = make_server(0); self.base = f'http://127.0.0.1:{server.server_port}'
            threading.Thread(target=server.serve_forever, daemon=True).start()
            try:
                self.assertEqual(self.call('GET', '/pack?job_id=fixture', token='wrong')[0], 401)
                self.assertEqual(self.call('GET', '/pack?job_id=nope')[0], 404)
                import tempfile
                pack = build(j, profile(), Path(paths.data_dir()) / 'packs' / 'fixture'); s.draft(j, pack)
                status, body = self.call('GET', '/pack?job_id=fixture')
                self.assertEqual(status, 409); self.assertIn('Application is not ready', body['blockers'])
                pack.update(adapter=fixture_adapter(), human_reviewed=True)
                s.db.execute('UPDATE applications SET payload=? WHERE job_id=?', (json.dumps(pack), 'fixture')); s.db.commit(); s.transition('fixture', 'ready')
                status, body = self.call('GET', '/pack?job_id=fixture')
                self.assertEqual(status, 200, body); self.assertTrue(body['gate_passed'])
                self.assertEqual(self.call('POST', '/attempt', {'job_id': 'fixture'})[0], 200)
                status, body = self.call('POST', '/result', {'job_id': 'fixture', 'status': 'submitted', 'receipt': ''})
                self.assertEqual(status, 409)  # no receipt, no "submitted"
                self.assertEqual(self.call('POST', '/result', {'job_id': 'fixture', 'status': 'uncertain'})[0], 200)
                self.assertEqual(self.call('POST', '/attempt', {'job_id': 'fixture'})[0], 409)  # uncertain is never re-attempted
            finally:
                server.shutdown(); server.server_close(); s.db.close()

    def test_refuses_a_short_token(self):
        from agent.server import make_server
        with patch.dict(os.environ, {'CAREER_AGENT_TOKEN': 'short'}):
            with self.assertRaises(SystemExit):
                make_server(0)


class Review(unittest.TestCase):
    def test_approve_retry_reconcile_and_no_auto_submit(self):
        from agent import review, cli
        with DataDir():
            cli.main(['seed'])
            s = Store(paths.db_path())
            with self.assertRaises(SystemExit) as e:
                review.main(['approve-pack', '--job', 'EXAMPLE-1', '--auto-submit', '--confirmed'])
            self.assertIn('click submit yourself', str(e.exception))
            with self.assertRaises(SystemExit):
                review.main(['verify-job', '--job', 'EXAMPLE-1'])  # no --confirmed
            job = next(j for j in s.jobs() if j['id'] == 'EXAMPLE-1')
            job['requirements'][3]['state'] = 'met'; job['requirements'][3]['evidence_ids'] = ['E03']; s.put_job(job)
            cli.main(['prepare', '--job', 'EXAMPLE-1'])
            review.main(['verify-job', '--job', 'EXAMPLE-1', '--confirmed'])
            adapter = Path(paths.data_dir()) / 'adapter.json'
            adapter.write_text(json.dumps(dict(fixture_adapter(), exact_url='https://example.org/jobs/1')))
            review.main(['approve-pack', '--job', 'EXAMPLE-1', '--adapter', str(adapter), '--confirmed'])
            self.assertEqual(s.applications()[0]['status'], 'ready')
            s.transition('EXAMPLE-1', 'submitting'); s.transition('EXAMPLE-1', 'uncertain')
            with self.assertRaises(ValueError):
                review.main(['retry', '--job', 'EXAMPLE-1', '--confirmed'])  # uncertain cannot be retried
            review.main(['reconcile', '--job', 'EXAMPLE-1', '--receipt', 'Ref 123 from employer email', '--confirmed'])
            app = s.applications()[0]
            self.assertEqual(app['status'], 'submitted'); self.assertEqual(app['payload']['receipt_kind'], 'manually_reconciled')
            s.db.close()

    def test_below_80_cannot_be_prepared(self):
        from agent import cli
        with DataDir():
            cli.main(['seed'])
            with self.assertRaises(SystemExit) as e:
                cli.main(['prepare', '--job', 'EXAMPLE-2'])
            self.assertIn('below 80%', str(e.exception))


if __name__ == '__main__':
    unittest.main()
